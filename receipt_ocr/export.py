import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [("file", "ファイル"), ("date", "日付"), ("store", "店名"), ("total", "合計金額"),
           ("category", "勘定科目"), ("needs_review", "要確認")]
_REVIEW_FILL = PatternFill("solid", start_color="FFF2CC")


def write_csv(rows: list[dict], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=[k for k, _ in COLUMNS], extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def write_xlsx(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "領収書"
    ws.append([label for _, label in COLUMNS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([("要確認" if k == "needs_review" and r[k] else "" if k == "needs_review" else r.get(k))
                   for k, _ in COLUMNS])
        if r["needs_review"] or r.get("category") == "要確認":
            for c in ws[ws.max_row]:
                c.fill = _REVIEW_FILL
    last = ws.max_row
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].number_format = '#,##0"円"'
    ws.append([])
    ws.append(["合計", None, None, f"=SUM(D2:D{last})"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.cell(ws.max_row, 4).number_format = '#,##0"円"'
    ws.cell(ws.max_row, 4).font = Font(bold=True)
    # 勘定科目別集計シート
    sm = wb.create_sheet("科目別集計")
    sm.append(["勘定科目", "合計金額"])
    for c in sm[1]:
        c.font = Font(bold=True)
    for acct in sorted({r.get("category") or "要確認" for r in rows}):
        i = sm.max_row + 1
        sm.append([acct, f'=SUMIF(領収書!E2:E{last},A{i},領収書!D2:D{last})'])
        sm.cell(i, 2).number_format = '#,##0"円"'
    for sheet in (ws, sm):
        for i, col in enumerate(sheet.columns, 1):
            width = max((len(str(c.value)) if c.value else 0) for c in col)
            sheet.column_dimensions[get_column_letter(i)].width = min(40, max(10, width * 1.6))
    ws.freeze_panes = "A2"
    wb.save(path)
