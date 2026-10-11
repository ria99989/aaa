import shutil
from pathlib import Path

import numpy as np
import pytest
from PIL import Image

from receipt_ocr.category import classify
from receipt_ocr.export import write_xlsx
from receipt_ocr.preprocess import estimate_skew, rotate

ROOT = Path(__file__).resolve().parent.parent


def test_classify():
    assert classify("紀伊國屋書店 新宿") == "新聞図書費"
    assert classify("JR東日本") == "旅費交通費"
    assert classify("不明な店") == "要確認"
    assert classify(None) == "要確認"


def test_skew_estimation_recovers_angle():
    gray = np.full((300, 600), 255, np.uint8)
    gray[100:130, 50:550] = 0
    gray[170:200, 50:550] = 0
    tilted = rotate(gray, 7.0)
    assert abs(abs(estimate_skew(tilted)) - 7.0) < 1.0


def test_xlsx_has_totals_and_summary(tmp_path):
    from openpyxl import load_workbook
    rows = [{"file": "a.png", "date": "2025-01-01", "store": "A", "total": 1000, "category": "会議費", "needs_review": False},
            {"file": "b.png", "date": None, "store": "B", "total": 500, "category": "要確認", "needs_review": True}]
    write_xlsx(rows, tmp_path / "o.xlsx")
    wb = load_workbook(tmp_path / "o.xlsx")
    assert wb.sheetnames == ["領収書", "科目別集計"]
    assert wb["領収書"]["D5"].value == "=SUM(D2:D3)"


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract 未インストール")
def test_ocr_survives_rotation_and_shadow(tmp_path):
    import cv2
    from receipt_ocr.engine import run
    im = cv2.imread(str(ROOT / "samples" / "receipt1.png")).astype(np.float32)
    h, w = im.shape[:2]
    im = cv2.warpAffine(im, cv2.getRotationMatrix2D((w / 2, h / 2), 6, 1), (w, h), borderValue=(255, 255, 255))
    im *= np.tile(np.linspace(0.45, 1.0, w, dtype=np.float32), (h, 1))[..., None]
    p = tmp_path / "bad.png"
    cv2.imwrite(str(p), np.clip(im, 0, 255).astype(np.uint8))
    row = run(p, tmp_path / "o.csv")[0]
    assert row["date"] == "2025-03-14" and row["total"] == 2180
